"""
Smoke test for docx_service.py (Stage 2).

Creates a synthetic .docx, exercises extract_text / extract_paragraphs /
write_comments, and saves the output to /tmp/test_output.docx so you can
open it in Word to verify the comments appear correctly.

Usage:
    python3 backend/scripts/test_docx_service.py
"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from docx import Document
from backend.services.docx_service import (
    extract_text,
    extract_paragraphs,
    extract_existing_comments,
    write_comments,
)


# ── Build a synthetic test doc ────────────────────────────────────────────────

def make_test_docx() -> bytes:
    doc = Document()
    doc.add_heading("How AI is Transforming Drug Development", level=1)
    doc.add_paragraph(
        "Pharmaceutical companies are increasingly leveraging artificial intelligence "
        "to accelerate the drug development lifecycle."
    )
    doc.add_heading("Key Benefits", level=2)
    doc.add_paragraph(
        "Process automation reduces manual errors and speeds up regulatory submissions "
        "by a significant margin."
    )
    doc.add_paragraph(
        "Appian's low-code platform enables teams to build scalable workflows without "
        "heavy IT involvement."
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def check(label, fn):
    print(f"\n{'='*60}")
    print(f"TEST: {label}")
    print("=" * 60)
    try:
        result = fn()
        print("✅ PASS")
        return result
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"❌ FAIL: {e}")
        return None


# ── Tests ─────────────────────────────────────────────────────────────────────

def test_extract_text(docx_bytes):
    text = extract_text(docx_bytes)
    assert "AI is Transforming" in text
    assert "Appian" in text
    print(f"   Extracted {len(text)} chars, {text.count(chr(10))+1} lines")
    print(f"   Preview: {text[:120]}...")
    return text


def test_extract_paragraphs(docx_bytes):
    paras = extract_paragraphs(docx_bytes)
    assert len(paras) >= 3
    print(f"   Found {len(paras)} non-empty paragraphs:")
    for p in paras:
        print(f"     [{p['style']:20s}] {p['text'][:60]}")
    return paras


def test_no_existing_comments(docx_bytes):
    comments = extract_existing_comments(docx_bytes)
    print(f"   Existing comments: {len(comments)} (expected 0)")
    assert len(comments) == 0
    return comments


def test_write_comments(docx_bytes):
    suggestions = [
        {
            "original_text": "Pharmaceutical companies are increasingly leveraging artificial intelligence",
            "suggestion": "Life sciences leaders are rapidly adopting AI",
            "rationale": "More specific to our ICP; 'life sciences leaders' signals authority.",
        },
        {
            "original_text": "Appian's low-code platform enables teams to build scalable workflows",
            "suggestion": "Appian's low-code platform empowers teams to rapidly deploy scalable workflows",
            "rationale": "Add 'rapidly deploy' — reinforces speed-to-value, our key differentiator.",
        },
    ]
    output = write_comments(docx_bytes, suggestions)
    assert len(output) > len(docx_bytes), "Output should be larger than input"
    print(f"   Input:  {len(docx_bytes):,} bytes")
    print(f"   Output: {len(output):,} bytes (+{len(output)-len(docx_bytes):,})")
    return output


def test_comments_readable(output_bytes):
    comments = extract_existing_comments(output_bytes)
    print(f"   Comments written and readable: {len(comments)}")
    for c in comments:
        print(f"     author={c['author']}")
        print(f"     text preview: {c['text'][:80]}...")
    assert len(comments) == 2, f"Expected 2 comments, got {len(comments)}"
    assert all(c["author"] == "BlogReviewAgent" for c in comments)


def test_save_for_manual_inspection(output_bytes):
    out_path = Path("/tmp/test_output.docx")
    out_path.write_bytes(output_bytes)
    print(f"   Saved to {out_path}")
    print(f"   Open in Word to verify 2 comments appear on the correct paragraphs.")


if __name__ == "__main__":
    print("\nBlogReviewAgent — docx_service Smoke Test")

    docx_bytes = make_test_docx()
    print(f"\nCreated test .docx ({len(docx_bytes):,} bytes)")

    check("1. extract_text", lambda: test_extract_text(docx_bytes))
    check("2. extract_paragraphs", lambda: test_extract_paragraphs(docx_bytes))
    check("3. no existing comments on fresh doc", lambda: test_no_existing_comments(docx_bytes))
    output = check("4. write_comments", lambda: test_write_comments(docx_bytes))
    if output:
        check("5. comments readable after write", lambda: test_comments_readable(output))
        check("6. save for manual Word inspection", lambda: test_save_for_manual_inspection(output))

    print(f"\n{'='*60}")
    print("Done.")
